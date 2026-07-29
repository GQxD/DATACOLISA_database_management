from pathlib import Path
from openpyxl import load_workbook

path = Path(r'c:/Users/anaubin/Desktop/ECAILLE BASE ORDINATEUR/colisa en cours/COLISA en cours.xlsx')
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb['Feuil1']
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
print('headers count', len(headers))
for name in ['Ecailles brutes', 'Montées', 'Empreintes', 'Otolithes']:
    print(name, [(i + 1, h) for i, h in enumerate(headers) if h and str(h).strip().lower() == name.lower()])

count_present = {'ecailles_brutes': 0, 'montées': 0, 'empreintes': 0, 'otolithes': 0}
count_notempty = {'ecailles_brutes': 0, 'montées': 0, 'empreintes': 0, 'otolithes': 0}
count_values = {'ecailles_brutes': {}, 'montées': {}, 'empreintes': {}, 'otolithes': {}}
for row in ws.iter_rows(min_row=2, values_only=True):
    for key, col_index in [('ecailles_brutes', 37), ('montées', 38), ('empreintes', 39), ('otolithes', 40)]:
        if len(row) > col_index and row[col_index] is not None:
            val = str(row[col_index]).strip()
            if val != '':
                count_notempty[key] += 1
                count_values[key][val] = count_values[key].get(val, 0) + 1
                if val.upper() not in {'NON', 'NO', 'FALSE', 'FAUX', 'N', '0'}:
                    count_present[key] += 1

print('count_notempty', count_notempty)
print('count_present', count_present)
for key in count_values:
    print('\n', key, 'examples', list(count_values[key].items())[:10])
wb.close()
