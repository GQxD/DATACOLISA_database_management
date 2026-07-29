from pathlib import Path
from openpyxl import load_workbook

path = Path(r'c:/Users/anaubin/Desktop/ECAILLE BASE ORDINATEUR/colisa en cours/COLISA en cours.xlsx')
print('path', path)
print('exists', path.exists())
if not path.exists():
    raise FileNotFoundError(path)
wb = load_workbook(path, read_only=True, data_only=True)
print('sheets', wb.sheetnames)
ws = wb[wb.sheetnames[0]]
print('max_row', ws.max_row, 'max_col', ws.max_column)
headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
print('headers', headers)
print('first rows:')
for r in range(2, min(12, ws.max_row + 1)):
    print(r, [ws.cell(r, c).value for c in range(1, ws.max_column + 1)])
wb.close()
