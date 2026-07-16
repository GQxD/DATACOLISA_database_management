import sys
from pathlib import Path
from openpyxl import load_workbook

# Ensure project import path
p = Path(__file__).resolve().parents[1]
import sys as _sys
_sys.path.insert(0, str(p))

from infrastructure.internal_target_workbook import build_numero_identification_formula


def find_header_row_and_col(ws, header_name: str = 'CODE IDENTIFICATION'):
    # Search first 8 rows for the header cell
    target = header_name.strip().lower()
    for r in range(1, 9):
        for c in range(1, min(60, ws.max_column + 1)):
            v = ws.cell(r, c).value
            if v and isinstance(v, str) and v.strip().lower() == target:
                return r, c
    # fallback: header row 1, conventional column 18
    return 1, 18


def restore_formulas(src: Path, dst: Path | None = None):
    if dst is None:
        dst = src.with_name(src.stem + '_restored' + src.suffix)

    wb = load_workbook(src, data_only=False)
    ws = wb.active

    header_row, code_col = find_header_row_and_col(ws)
    print(f'Found header at row {header_row}, column {code_col}')

    changed = 0
    for r in range(header_row + 1, ws.max_row + 1):
        # skip completely empty rows
        has_data = any(ws.cell(r, c).value not in (None, '') for c in range(1, ws.max_column + 1))
        if not has_data:
            continue

        cell = ws.cell(r, code_col)
        val = cell.value
        needs = False
        if val is None or (isinstance(val, str) and not val.startswith('=')):
            # reinsert formula
            formula = build_numero_identification_formula(r)
            if formula:
                cell.value = formula
                changed += 1

    if changed == 0:
        print('No cells updated (no empty or non-formula code identification cells found).')
    else:
        print(f'Inserted formulas into {changed} cells.')

    wb.save(dst)
    wb.close()
    print('Wrote', dst)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python restore_eid_formulas.py "path/to/COLISA en cours.xlsx" [output.xlsx]')
        raise SystemExit(1)
    src = Path(sys.argv[1])
    if not src.exists():
        print('File not found:', src)
        raise SystemExit(2)
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    restore_formulas(src, dst)
