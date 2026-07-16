from openpyxl import load_workbook
from pathlib import Path
p = Path(__file__).resolve().parents[1] / 'assets' / 'colisa_logiciel_template.xlsx'
wb = load_workbook(p, data_only=False)
ws = wb.active
out = []
out.append(f'sheet: {ws.title}')
count = 0
for r in range(1, 30):
    for c in range(1, 40):
        cell = ws.cell(r, c)
        v = cell.value
        t = cell.data_type
        has_f = isinstance(v, str) and v.startswith('=')
        if has_f or t == 'f':
            out.append(f"{cell.coordinate} | value={v!r} | data_type={t}")
            count += 1
out.append(f'found: {count}')
wb.close()

dest = Path(__file__).resolve().parents[0] / 'inspect_output_utf8.txt'
dest.write_text('\n'.join(out), encoding='utf-8')
print('WROTE', dest)
