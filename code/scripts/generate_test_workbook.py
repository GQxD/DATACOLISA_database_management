import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from infrastructure.internal_target_workbook import create_internal_target_workbook, build_numero_identification_formula
from openpyxl import load_workbook
import datetime as dt

out = Path(__file__).resolve().parents[0] / 'test_out.xlsx'
# Create base workbook
create_internal_target_workbook(out, __import__('openpyxl'))

wb = load_workbook(out)
ws = wb.active
# Fill required columns for rows 2 and 3
# Column indices: E=5, F=6, K=11, L=12, S=19
ws.cell(2, 12).value = 'LEMAN'  # L2
ws.cell(2, 6).value = 'A'       # F2 (espece)
ws.cell(2, 11).value = dt.datetime(2020,1,2)  # K2 date
ws.cell(2, 19).value = '1'     # S2 numero individu
ws.cell(2, 5).value = 'T00001' # E2 code echantillon
# Set formula
ws.cell(2, 18).value = build_numero_identification_formula(2)

# Row 3
ws.cell(3, 12).value = 'LEMAN'
ws.cell(3, 6).value = 'A'
ws.cell(3, 11).value = dt.datetime(2020,1,3)
ws.cell(3, 19).value = '2'
ws.cell(3, 5).value = 'T00002'
ws.cell(3, 18).value = build_numero_identification_formula(3)

wb.save(out)
wb.close()
print('WROTE', out)

# Re-open and show formulas stored
wb = load_workbook(out, data_only=False)
ws = wb.active
print('Row2 EID formula:', ws.cell(2,18).value)
print('Row3 EID formula:', ws.cell(3,18).value)
wb.close()
